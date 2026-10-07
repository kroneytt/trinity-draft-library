import openpyxl

file_path = "data/cards.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb.active

header = {cell.value: i for i, cell in enumerate(ws[1])}

# Let's see what the sub type column is actually called
sub_type_col = header.get("Sub-Type") or header.get("Sub_Type") or header.get("Subtype")

for row in ws.iter_rows(min_row=2):
    epic_cell = row[header["Epic_Rule_or_Burst"]]
    if epic_cell.value and "add 1 chosen type from outside the game to your hand" in str(epic_cell.value):
        epic_cell.value = str(epic_cell.value).replace(
            "add 1 chosen type from outside the game to your hand",
            "At the start of the game, add your chosen Epic to your hand"
        )
    
    type_cell = row[header["Type"]]
    if sub_type_col is not None:
        sub_cell = row[sub_type_col]
        if type_cell.value == "Spell" and sub_cell.value == "Set":
            sub_cell.value = "Installation"
        if type_cell.value == "Spell" and sub_cell.value == "Equipment":
            sub_cell.value = "Equip"
            
    eff_cell = row[header["Effects"]]
    if eff_cell.value:
        eff_cell.value = str(eff_cell.value).replace("[Set]", "[Installation]")
        eff_cell.value = str(eff_cell.value).replace("[Equipment]", "[Equip]")
        eff_cell.value = str(eff_cell.value).replace("Traitor", "Betrayal")

    if epic_cell.value and "Traitor" in str(epic_cell.value):
        epic_cell.value = str(epic_cell.value).replace("Traitor", "Betrayal")

wb.save(file_path)
print("Updated cards.xlsx")
